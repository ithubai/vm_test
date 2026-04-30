"""Genera un file Excel a partire dalle righe d'ordine estratte."""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


_INTESTAZIONI = ["Codice Articolo", "Descrizione", "Quantità", "Unità di Misura"]
_CHIAVI =       ["codice_articolo", "descrizione",  "quantita", "unita_misura"]

_COLORE_HEADER = "1F4E79"   # blu scuro
_COLORE_TESTO  = "FFFFFF"   # bianco
_COLORE_RIGA_ALT = "D6E4F0" # azzurro chiaro per righe alternate


def _bordo_sottile() -> Border:
    lato = Side(style="thin", color="AAAAAA")
    return Border(left=lato, right=lato, top=lato, bottom=lato)


def scrivi_excel(righe: list[dict], percorso_output: str) -> str:
    """
    Scrive le righe d'ordine in un file Excel formattato.

    Parametri
    ---------
    righe : list[dict]
        Lista di dizionari con le chiavi codice_articolo, descrizione,
        quantita, unita_misura.
    percorso_output : str
        Percorso del file .xlsx da creare.

    Ritorna
    -------
    str
        Percorso assoluto del file creato.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ordine"

    # --- Intestazioni ---
    font_header = Font(bold=True, color=_COLORE_TESTO, size=11)
    fill_header = PatternFill("solid", fgColor=_COLORE_HEADER)
    allineamento_centro = Alignment(horizontal="center", vertical="center")

    for col_idx, intestazione in enumerate(_INTESTAZIONI, start=1):
        cella = ws.cell(row=1, column=col_idx, value=intestazione)
        cella.font = font_header
        cella.fill = fill_header
        cella.alignment = allineamento_centro
        cella.border = _bordo_sottile()

    # --- Righe dati ---
    fill_alt = PatternFill("solid", fgColor=_COLORE_RIGA_ALT)

    for row_idx, riga in enumerate(righe, start=2):
        usa_fill_alt = (row_idx % 2 == 0)
        for col_idx, chiave in enumerate(_CHIAVI, start=1):
            valore = riga.get(chiave, "")
            cella = ws.cell(row=row_idx, column=col_idx, value=valore)
            cella.border = _bordo_sottile()
            if usa_fill_alt:
                cella.fill = fill_alt
            # Quantità allineata a destra
            if chiave == "quantita":
                cella.alignment = Alignment(horizontal="right")

    # --- Larghezze colonne auto ---
    larghezze_min = [16, 40, 12, 16]
    for col_idx, larghezza_min in enumerate(larghezze_min, start=1):
        max_len = larghezza_min
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cella in row:
                if cella.value:
                    max_len = max(max_len, len(str(cella.value)) + 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len

    # Altezza intestazione
    ws.row_dimensions[1].height = 20

    # Congela la riga di intestazione
    ws.freeze_panes = "A2"

    # --- Contatore in fondo ---
    riga_totale = len(righe) + 2
    ws.cell(row=riga_totale, column=1, value="Totale righe:").font = Font(bold=True)
    ws.cell(row=riga_totale, column=2, value=len(righe)).font = Font(bold=True)

    Path(percorso_output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(percorso_output)
    return str(Path(percorso_output).resolve())
